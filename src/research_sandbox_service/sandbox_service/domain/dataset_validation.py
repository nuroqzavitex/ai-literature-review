"""Fail-closed upload validation before a dataset reaches object storage."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import PurePath
import re
from typing import Literal
import zipfile

from openpyxl import load_workbook
import pyarrow.parquet as pq

from sandbox_service.domain.datasets import DatasetClassification
from sandbox_service.domain.errors import DatasetClassificationRejected, DatasetTooLarge, DatasetTypeNotAllowed, DatasetValidationFailed


DatasetFormat = Literal["csv", "xlsx", "parquet"]


@dataclass(frozen=True)
class ValidatedDataset:
    format: DatasetFormat
    media_type: str
    column_names: list[str]
    row_count: int


class DatasetValidator:
    MAX_BYTES = 50 * 1024 * 1024
    MAX_ROWS = 1_000_000
    MAX_COLUMNS = 1_000
    MAX_XLSX_UNCOMPRESSED_BYTES = 500 * 1024 * 1024
    MAX_XLSX_EXPANSION_RATIO = 100

    _MIME_TYPES: dict[DatasetFormat, set[str]] = {
        "csv": {"text/csv", "application/csv", "text/plain"},
        "xlsx": {
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/octet-stream",
        },
        "parquet": {"application/vnd.apache.parquet", "application/parquet", "application/octet-stream"},
    }
    _FORMULA_PREFIXES = ("=", "+", "-", "@")

    def validate(
        self,
        *,
        filename: str,
        data: bytes,
        declared_media_type: str | None,
        classification: DatasetClassification,
    ) -> ValidatedDataset:
        if classification is not DatasetClassification.NON_SENSITIVE:
            raise DatasetClassificationRejected("Only non_sensitive datasets are accepted in S2")
        if not data:
            raise DatasetValidationFailed("Dataset is empty")
        if len(data) > self.MAX_BYTES:
            raise DatasetTooLarge("Dataset exceeds the 50 MB upload limit")
        if PurePath(filename).name != filename or not filename.strip():
            raise DatasetValidationFailed("Dataset filename must not contain a path")

        suffix = PurePath(filename).suffix.lower()
        detected = self._detect_format(data)
        expected_suffix = f".{detected}"
        if suffix != expected_suffix:
            raise DatasetTypeNotAllowed("File extension does not match the detected dataset format")
        if declared_media_type and declared_media_type.lower() not in self._MIME_TYPES[detected]:
            raise DatasetTypeNotAllowed("Declared MIME type does not match the detected dataset format")

        if detected == "csv":
            columns, rows = self._validate_csv(data)
        elif detected == "xlsx":
            columns, rows = self._validate_xlsx(data)
        else:
            columns, rows = self._validate_parquet(data)
        self._validate_columns(columns)
        if rows > self.MAX_ROWS or len(columns) > self.MAX_COLUMNS:
            raise DatasetValidationFailed("Dataset exceeds configured row or column limits")
        return ValidatedDataset(format=detected, media_type=self._canonical_media_type(detected), column_names=columns, row_count=rows)

    def _detect_format(self, data: bytes) -> DatasetFormat:
        if len(data) >= 8 and data[:4] == b"PAR1" and data[-4:] == b"PAR1":
            return "parquet"
        if data[:4] == b"PK\x03\x04":
            self._inspect_xlsx_zip(data)
            return "xlsx"
        try:
            data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DatasetTypeNotAllowed("File magic is not CSV, XLSX, or Parquet") from exc
        if b"\x00" in data[:4096]:
            raise DatasetTypeNotAllowed("Text dataset contains binary bytes")
        return "csv"

    def _validate_csv(self, data: bytes) -> tuple[list[str], int]:
        text = data.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(StringIO(text), dialect)
        try:
            columns = next(reader)
        except StopIteration as exc:
            raise DatasetValidationFailed("CSV must contain a header row") from exc
        rows = 0
        for row in reader:
            rows += 1
            if any(self._is_formula(value) for value in row):
                raise DatasetValidationFailed("Formula-like CSV values are not accepted")
        return columns, rows

    def _validate_xlsx(self, data: bytes) -> tuple[list[str], int]:
        self._inspect_xlsx_zip(data)
        try:
            workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_vba=False)
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=False)
            header_cells = next(rows)
            columns = ["" if cell.value is None else str(cell.value) for cell in header_cells]
            count = 0
            for row in rows:
                count += 1
                for cell in row:
                    value = cell.value
                    if cell.data_type == "f" or (isinstance(value, str) and self._is_formula(value)):
                        raise DatasetValidationFailed("Formula-like XLSX values are not accepted")
            workbook.close()
            return columns, count
        except DatasetValidationFailed:
            raise
        except Exception as exc:
            raise DatasetValidationFailed("XLSX parser rejected the upload") from exc

    def _validate_parquet(self, data: bytes) -> tuple[list[str], int]:
        try:
            parquet_file = pq.ParquetFile(BytesIO(data))
            metadata = parquet_file.metadata
            if metadata is None:
                raise ValueError("missing Parquet metadata")
            uncompressed = sum(
                column.total_uncompressed_size
                for row_group_index in range(metadata.num_row_groups)
                for column in (
                    metadata.row_group(row_group_index).column(column_index)
                    for column_index in range(metadata.row_group(row_group_index).num_columns)
                )
            )
            if uncompressed > self.MAX_XLSX_UNCOMPRESSED_BYTES or uncompressed / max(len(data), 1) > self.MAX_XLSX_EXPANSION_RATIO:
                raise DatasetValidationFailed("Parquet compression exceeds decompression-bomb limits")
            schema = parquet_file.schema_arrow
            return list(schema.names), metadata.num_rows
        except Exception as exc:
            raise DatasetValidationFailed("Parquet parser rejected the upload") from exc

    def _inspect_xlsx_zip(self, data: bytes) -> None:
        try:
            with zipfile.ZipFile(BytesIO(data)) as archive:
                names = set(archive.namelist())
                if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                    raise DatasetTypeNotAllowed("ZIP upload is not an XLSX workbook")
                if any(name.lower().endswith("vbaProject.bin".lower()) for name in names):
                    raise DatasetTypeNotAllowed("Macro-enabled workbooks are not accepted")
                uncompressed = sum(item.file_size for item in archive.infolist())
                compressed = max(sum(item.compress_size for item in archive.infolist()), 1)
                if uncompressed > self.MAX_XLSX_UNCOMPRESSED_BYTES or uncompressed / compressed > self.MAX_XLSX_EXPANSION_RATIO:
                    raise DatasetValidationFailed("XLSX archive exceeds compression-bomb limits")
        except DatasetTypeNotAllowed:
            raise
        except DatasetValidationFailed:
            raise
        except zipfile.BadZipFile as exc:
            raise DatasetTypeNotAllowed("ZIP signature is not a valid XLSX workbook") from exc

    def _validate_columns(self, columns: list[str]) -> None:
        if not columns or any(not value.strip() for value in columns):
            raise DatasetValidationFailed("Dataset must contain non-empty column names")
        normalized: set[str] = set()
        for name in columns:
            if len(name) > 128 or any(ord(character) < 32 for character in name):
                raise DatasetValidationFailed("Dataset has an unsafe column name")
            if self._is_formula(name) or not re.match(r"^[^./\\:]+$", name):
                raise DatasetValidationFailed("Dataset has an unsafe column name")
            key = name.casefold().strip()
            if key in normalized:
                raise DatasetValidationFailed("Dataset has duplicate column names")
            normalized.add(key)

    def _is_formula(self, value: str) -> bool:
        stripped = value.lstrip()
        # Preserve legitimate negative numeric observations while rejecting a
        # spreadsheet expression such as "-CMD(...)".
        if re.fullmatch(r"-\d+(?:\.\d+)?(?:[eE][+-]?\d+)?", stripped):
            return False
        return stripped.startswith(self._FORMULA_PREFIXES)

    @staticmethod
    def _canonical_media_type(format_: DatasetFormat) -> str:
        return {
            "csv": "text/csv",
            "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "parquet": "application/vnd.apache.parquet",
        }[format_]
